from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

from .logger import get_logger
from .paths import sample_data_dir, source_path, mpbse_data_dir, ROOT, load_json


log = get_logger("sample_importer")


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(parts)


def import_sample_exam() -> dict:
    path = source_path("sample_exam_docx")
    if not path.exists():
        raise FileNotFoundError(f"Sample exam not found: {path}")
    text = _docx_text(path)
    out_dir = sample_data_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "grade10_math_2026_sample.txt").write_text(text, encoding="utf-8")

    chapters_mentioned = []
    for line in text.splitlines()[:12]:
        m = re.match(r"^\s*\d+\s+(.+?):?\s*$", line)
        if m and "Real" in m.group(1) or "Polynomial" in line or "Linear" in line or "Quadratic" in line:
            chapters_mentioned.append(line.strip())

    record = {
        "sample_exam_id": "grade10_math_2026_sample_qe",
        "title": "10th Math 2026 Sample Exam",
        "source_file": str(path),
        "status": "READY",
        "grade": 10,
        "subject": "Mathematics",
        "stated_marks": 75,
        "stated_duration": "3 hours",
        "coverage_note": "Chapters listed on the paper: Real Numbers, Polynomials, Pair of Linear Equations in Two Variables, Quadratic Equations.",
        "style_role": "STYLE/EXAMPLE evidence only. Not higher authority than explicit MPBSE rules.",
        "observed_structure": {
            "Q1_Q5_objective": True,
            "Q6_Q17_2mark": True,
            "Q18_Q20_3mark": True,
            "Q21_plus_4mark": True,
            "extra_long_items_q24_q25": True,
            "internal_choice_note": "Long-answer heading says 'कोई 3' (any 3). Not stored as an official MPBSE Class 10 Math rule."
        },
        "header_language_in_sample": "hi",
        "body_language_in_sample": "hi",
        "warnings": [
            "Sample paper includes Q24 and Q25 beyond the Q21-Q23 pattern used in the Exam Profile.",
            "Some option/formula glyphs did not extract cleanly from the DOCX; used as style evidence only."
        ],
    }
    (out_dir / "grade10_math_2026_sample.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log.info("Imported sample exam %s (%s chars)", record["sample_exam_id"], len(text))
    return record


def import_mpbse_snapshot() -> dict:
    marking = source_path("mpbse_marking_docx")
    translation = source_path("mpbse_translation_docx")
    notes = load_json(ROOT / "config" / "source_notes.json")
    profile = load_json(ROOT / "config" / "exam_profiles" / "grade10_math.json")

    snapshot = {
        "source_files": {
            "marking_docx": str(marking) if marking.exists() else None,
            "translation_docx": str(translation) if translation.exists() else None,
            "exam_pattern_pdf": None,
        },
        "question_structure_status": "READY",
        "chapter_weightage_status": profile["chapter_weightage_status"],
        "class10_mathematics_subject_block": "not_listed_in_translation",
        "usable_75_mark_pattern": profile["question_structure"],
        "flags": notes["flags"],
        "explicit_non_use": [
            "Class 9 Mathematics chapter weightage table is present but MUST NOT be used as Class 10 weightage.",
            "Class 12 Mathematics is 80 marks with its own chapter table and MUST NOT be used as Class 10 weightage."
        ],
    }
    out = mpbse_data_dir()
    out.mkdir(parents=True, exist_ok=True)
    (out / "grade10_question_structure.json").write_text(
        json.dumps(snapshot, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    log.info(
        "MPBSE snapshot saved. question_structure=%s chapter_weightage=%s",
        snapshot["question_structure_status"],
        snapshot["chapter_weightage_status"],
    )
    return snapshot


if __name__ == "__main__":
    import_sample_exam()
    import_mpbse_snapshot()
