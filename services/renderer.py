from __future__ import annotations

import re
import struct
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .logger import get_logger
from .match_columns_render import (
    is_match_item,
    is_residual_match_item,
    mapping_lines,
    match_table_payload,
)


log = get_logger("renderer")

SUB_LABELS = {
    "i": "(i)",
    "ii": "(ii)",
    "iii": "(iii)",
    "iv": "(iv)",
    "v": "(v)",
    "vi": "(vi)",
}

# Central presentation styles — reusable across subjects. Sizes in PDF points.
STYLES = {
    "school": {"size": 15.0, "bold": True, "leading": 1.15},
    "title": {"size": 14.5, "bold": True, "leading": 1.15},
    "meta": {"size": 11.5, "bold": False, "leading": 1.15},
    "section": {"size": 13.5, "bold": True, "leading": 1.15},
    "label": {"size": 12.5, "bold": True, "leading": 1.15},
    "question": {"size": 12.0, "bold": True, "leading": 1.15},
    "body": {"size": 11.5, "bold": False, "leading": 1.15},
    "subq": {"size": 11.5, "bold": False, "leading": 1.15},
    "option": {"size": 11.5, "bold": False, "leading": 1.15},
    "passage": {"size": 11.5, "bold": False, "leading": 1.18},
    "or_line": {"size": 12.0, "bold": True, "leading": 1.15},
    "footer": {"size": 9.5, "bold": False, "leading": 1.1},
}

MARGIN_PT = 54.0  # 0.75 inch
SPACE_BEFORE_QUESTION = 10.0
SPACE_AFTER_QUESTION = 5.0
SPACE_BETWEEN_SUBQ = 5.0
SPACE_GRAMMAR_ITEM = 4.0
SPACE_PASSAGE = 6.0
SPACE_OR = 8.0
INDENT_SUBQ = 18.0
INDENT_OPTION = 36.0
INDENT_PASSAGE = 16.0
KEEP_MAX_PT = 230.0
EMPHASIS_RE = re.compile(
    r"\b(?:Any\s+(?:ONE|TWO|THREE|FOUR|one|two|three|four|\d+)|OR)\b",
    re.I,
)
SECTION_INS_RE = re.compile(
    r"Section\s+([A-Z])\s*:\s*([^(]+?)\s*(?:\((\d+)\s*marks?\))?",
    re.I,
)


@dataclass
class DocLine:
    text: str
    style: str = "body"
    indent: float = 0.0
    txt_indent: int = 0
    center: bool = False
    space_before: float = 0.0
    space_after: float = 0.0
    keep_with_next: bool = False
    group: str = ""
    image_path: str = ""
    image_height: float = 0.0
    table: dict | None = None


def _syllabus_lines(exam: dict) -> list[str]:
    from .content_gate import syllabus_display_from_exam

    display = syllabus_display_from_exam(exam)
    lines = ["SYLLABUS"]
    for group in display.get("groups") or []:
        book = (group.get("book") or "").strip()
        items = [i for i in (group.get("items") or []) if str(i).strip()]
        if book:
            lines.append(f"{book}:")
        for item in items:
            lines.append(f"• {item}")
    lines.append("")
    return lines


def _item_question_text(item: dict) -> str:
    return (item.get("question_text") or item.get("question_text_hindi") or "").strip()


def _item_answer_text(item: dict) -> str:
    return (item.get("answer_key") or item.get("answer_key_hindi") or "").strip()


def _format_option(opt, *, letter_style: bool = False) -> str:
    from .english_print_cleanup import format_option_line

    return format_option_line(opt, letter_style=letter_style)


def _sub_display(sub: str | None) -> str:
    if not sub:
        return ""
    raw = str(sub)
    if "-" in raw:
        raw = raw.split("-", 1)[-1]
    return SUB_LABELS.get(raw, f"({raw})")


def _norm_section_class(text: str) -> str:
    token = re.split(r"\s+", (text or "").strip())[0] if text else ""
    return token.upper()


def _sections_from_exam(exam: dict) -> list[dict]:
    if exam.get("sections"):
        out = []
        for raw in exam["sections"]:
            rec = dict(raw)
            rec["content_class"] = _norm_section_class(
                rec.get("content_class") or rec.get("name") or ""
            )
            out.append(rec)
        return out
    secs = []
    for ins in (exam.get("header") or {}).get("instructions") or []:
        m = SECTION_INS_RE.search(str(ins))
        if not m:
            continue
        name = m.group(2).strip().rstrip(".")
        secs.append(
            {
                "section_id": m.group(1).upper(),
                "name": name,
                "marks": int(m.group(3)) if m.group(3) else None,
                "content_class": _norm_section_class(name),
            }
        )
    return secs


def _section_banner(sec: dict) -> str:
    sid = str(sec.get("section_id") or "").strip()
    name = (sec.get("name") or "").strip()
    cls = (sec.get("content_class") or "").strip()
    if name and len(name.split()) <= 2:
        display = name.upper()
    elif cls:
        display = cls.replace("_", " ").upper()
    else:
        display = name.upper()
    if sid:
        banner = f"SECTION {sid} — {display}"
    else:
        banner = display
    marks = sec.get("marks")
    if marks is not None and str(marks) != "":
        banner += f"  ({marks} Marks)"
    return banner


def _poetry_doc_lines(passage: str, group: str) -> list[DocLine]:
    out: list[DocLine] = []
    rows = (passage or "").split("\n")
    out.append(DocLine("", "passage", space_before=SPACE_PASSAGE, group=group, keep_with_next=True))
    for i, raw in enumerate(rows):
        line = raw.rstrip()
        is_title = i == 0 and bool(line.strip()) and not line.strip().startswith("—")
        is_author = line.strip().startswith("—")
        is_blank = not line.strip()
        out.append(
            DocLine(
                line,
                "label" if is_title else "passage",
                indent=0 if is_title or is_author else INDENT_PASSAGE + 10,
                txt_indent=0 if is_title or is_author else 2,
                center=is_title or is_author,
                space_before=6 if is_blank else (2 if is_title else 0),
                space_after=8 if is_author else (8 if is_title else 0),
                keep_with_next=True,
                group=group,
            )
        )
    if out:
        out[-1].space_after = max(out[-1].space_after, SPACE_PASSAGE)
    return out


def _q_label(item: dict) -> str:
    qn = item.get("question_number")
    slot = item.get("sub_slot")
    label = f"Q{qn}"
    if slot:
        label += f"({slot})"
    return label


def _attempt_extra(item: dict, heading: str) -> str:
    any_n = item.get("attempt_any")
    if any_n and not re.search(r"\bany\s+\w+\b", str(heading), re.I):
        return f"  (any {any_n})"
    return ""


def _is_english(exam: dict) -> bool:
    h = exam.get("header") or {}
    subject = str(h.get("subject") or exam.get("subject") or "Mathematics")
    return subject.lower() == "english"


def _prepared_items(exam: dict) -> list[dict]:
    items = list(exam.get("questions") or [])
    if _is_english(exam):
        from .english_print_cleanup import sanitize_item

        items = [sanitize_item(q) for q in items]
    return items


def _default_instructions(exam: dict) -> list[str]:
    h = exam.get("header") or {}
    return list(
        h.get("instructions")
        or [
            "All questions are compulsory.",
            "Questions 1 to 5 are objective (1 mark each sub-question).",
            "Questions 6 to 17 are very short answer (2 marks each).",
            "Questions 18 to 20 are short answer (3 marks each).",
            "Questions 21 to 23 are long answer (4 marks each).",
        ]
    )


def build_exam_document(exam: dict) -> list[DocLine]:
    h = exam.get("header") or {}
    items = _prepared_items(exam)
    is_english = _is_english(exam)
    sections = _sections_from_exam(exam)
    section_by_class = {s.get("content_class"): s for s in sections if s.get("content_class")}
    lines: list[DocLine] = []

    school = h.get("school_line") or "School Examination"
    title = h.get("exam_title") or exam.get("exam_name") or "Mathematics Examination"
    grade = h.get("grade", exam.get("grade", 10))
    subject = h.get("subject", exam.get("subject", "Mathematics"))
    marks = h.get("total_marks", exam.get("total_marks"))
    minutes = h.get("duration_minutes", exam.get("duration_minutes"))

    lines.append(DocLine(str(school), "school", space_after=2))
    lines.append(DocLine(str(title), "title", space_after=6))
    lines.append(DocLine(f"Grade: {grade}", "meta", keep_with_next=True, group="header-meta"))
    lines.append(DocLine(f"Subject: {subject}", "meta", keep_with_next=True, group="header-meta"))
    lines.append(DocLine(f"Time: {minutes} minutes", "meta", keep_with_next=True, group="header-meta"))
    lines.append(DocLine(f"Maximum Marks: {marks}", "meta", space_after=10, group="header-meta"))

    lines.append(DocLine("SYLLABUS", "label", space_before=4, space_after=4, keep_with_next=True, group="syllabus"))
    for raw in _syllabus_lines(exam)[1:]:
        if not raw.strip():
            continue
        lines.append(DocLine(raw, "body", indent=8, txt_indent=0, group="syllabus", space_after=1))
    lines[-1].space_after = 8

    lines.append(
        DocLine("INSTRUCTIONS", "label", space_before=2, space_after=4, keep_with_next=True, group="instructions")
    )
    for ins in _default_instructions(exam):
        lines.append(
            DocLine(
                str(ins),
                "body",
                indent=14,
                txt_indent=2,
                group="instructions",
                space_after=2,
            )
        )
    lines[-1].space_after = 8

    current_q = None
    current_slot = None
    current_section = None
    last_passage = None
    lead_group = ""
    first_item = False
    for item in items:
        qn = item.get("question_number")
        slot = item.get("sub_slot")
        cls = (item.get("content_class") or "").upper()
        is_grammar = cls == "GRAMMAR"
        is_mcq = bool(item.get("options"))
        sec = section_by_class.get(cls)

        if sec and sec.get("section_id") != current_section:
            current_section = sec.get("section_id")
            banner = _section_banner(sec)
            lines.append(DocLine("", "body", space_before=4, space_after=2))
            lines.append(
                DocLine(
                    banner,
                    "section",
                    space_before=8,
                    space_after=8,
                    keep_with_next=True,
                    group=f"sec-{current_section}",
                )
            )

        new_main = qn != current_q or slot != current_slot
        if new_main:
            current_q = qn
            current_slot = slot
            last_passage = None
            heading = item.get("section_heading") or f"Question {qn}"
            extra = _attempt_extra(item, heading)
            qtext = f"{_q_label(item)}. {heading}{extra}".strip()
            lead_group = f"q{qn}{slot or ''}-lead"
            first_item = True
            lines.append(
                DocLine(
                    qtext,
                    "question",
                    space_before=SPACE_BEFORE_QUESTION,
                    space_after=SPACE_AFTER_QUESTION,
                    keep_with_next=True,
                    group=lead_group,
                )
            )

        if is_english:
            from .english_print_cleanup import printable_passage

            passage = printable_passage(item)
        else:
            passage = (item.get("passage") or "").strip()
        if passage and passage != last_passage:
            if (item.get("kind") or "").lower() == "poem" or "\n" in passage:
                lines.extend(_poetry_doc_lines(passage, lead_group or f"q{qn}{slot or ''}-pass"))
            else:
                lines.append(
                    DocLine(
                        passage,
                        "passage",
                        indent=INDENT_PASSAGE,
                        txt_indent=2,
                        space_before=SPACE_PASSAGE,
                        space_after=SPACE_PASSAGE,
                        keep_with_next=True,
                        group=lead_group or f"q{qn}{slot or ''}-pass",
                    )
                )
            last_passage = passage

        image_path = (item.get("picture_image") or "").strip()
        if item.get("format_id") == "PICTURE_COMPOSITION" and image_path:
            lines.append(
                DocLine(
                    "",
                    "body",
                    indent=INDENT_PASSAGE,
                    space_before=6,
                    space_after=6,
                    keep_with_next=True,
                    group=f"q{qn}-task",
                    image_path=image_path,
                    image_height=168.0,
                )
            )
        topics = item.get("essay_topics") or []
        if topics and item.get("format_id") == "ESSAY":
            for i, t in enumerate(topics, 1):
                lines.append(
                    DocLine(
                        f"({chr(96 + i)}) {t}",
                        "body",
                        indent=INDENT_SUBQ,
                        txt_indent=2,
                        space_after=3,
                        group=f"q{qn}-task",
                    )
                )

        sub = item.get("subquestion_number")
        label = _sub_display(sub)
        qbody = _item_question_text(item)
        sub_gap = SPACE_GRAMMAR_ITEM if is_grammar else SPACE_BETWEEN_SUBQ
        item_group = lead_group if first_item and lead_group else f"q{qn}{slot or ''}-{sub or 'main'}"
        first_item = False
        if is_match_item(item):
            if is_residual_match_item(item):
                continue
            payload = match_table_payload(item)
            if payload:
                table_group = f"q{qn}{slot or ''}-match"
                lines.append(
                    DocLine(
                        "",
                        "body",
                        indent=INDENT_SUBQ,
                        txt_indent=2,
                        space_before=sub_gap,
                        space_after=6,
                        keep_with_next=True,
                        group=table_group,
                        table={
                            "headers": ["Column A", "Column B"],
                            "rows": payload["rows"],
                            "col_weights": [0.62, 0.38],
                        },
                    )
                )
            continue
        skip_body = bool(item.get("suppress_question_text")) or (
            item.get("format_id") == "ESSAY" and item.get("essay_topics")
        ) or item.get("format_id") == "NOTE_MAKING"
        if qbody and not skip_body:
            prefix = f"{label} " if label else ""
            lines.append(
                DocLine(
                    prefix + qbody,
                    "subq" if label else "body",
                    indent=INDENT_SUBQ,
                    txt_indent=2,
                    space_before=sub_gap,
                    space_after=3,
                    keep_with_next=is_mcq,
                    group=item_group,
                )
            )
        if item.get("options"):
            for opt in item["options"]:
                lines.append(
                    DocLine(
                        _format_option(opt, letter_style=is_english),
                        "option",
                        indent=INDENT_OPTION,
                        txt_indent=6,
                        space_after=1.5,
                        group=item_group,
                    )
                )
            lines[-1].space_after = 4
        choice_b = (item.get("choice_b") or "").strip()
        if choice_b and item.get("internal_or"):
            or_group = f"q{qn}{slot or ''}-or"
            lines.append(
                DocLine(
                    "OR",
                    "or_line",
                    center=True,
                    space_before=SPACE_OR,
                    space_after=SPACE_OR,
                    keep_with_next=True,
                    group=or_group,
                )
            )
            lines.append(
                DocLine(
                    choice_b,
                    "body",
                    indent=INDENT_SUBQ,
                    txt_indent=2,
                    group=or_group,
                )
            )

    if is_english:
        collapsed: list[DocLine] = []
        for ln in lines:
            if collapsed and ln.text.strip() and ln.text.strip() == collapsed[-1].text.strip():
                continue
            collapsed.append(ln)
        lines = collapsed
    return lines


def build_answer_key_document(exam: dict) -> list[DocLine]:
    h = exam.get("header") or {}
    sections = _sections_from_exam(exam)
    section_by_class = {s.get("content_class"): s for s in sections if s.get("content_class")}
    lines = [
        DocLine("ANSWER KEY", "school", space_after=4),
        DocLine(str(h.get("exam_title") or exam.get("exam_name") or ""), "title", space_after=2),
        DocLine(
            f"Grade {h.get('grade', 10)} {h.get('subject', exam.get('subject', 'Mathematics'))}",
            "meta",
            space_after=10,
        ),
    ]
    current_section = None
    last_q = None
    for item in exam.get("questions") or []:
        cls = (item.get("content_class") or "").upper()
        sec = section_by_class.get(cls)
        if sec and sec.get("section_id") != current_section:
            current_section = sec.get("section_id")
            lines.append(
                DocLine(
                    _section_banner(sec),
                    "section",
                    space_before=10,
                    space_after=6,
                )
            )
        qn = item.get("question_number")
        sub = item.get("subquestion_number")
        slot = item.get("sub_slot")
        if is_match_item(item) and is_residual_match_item(item):
            continue
        label = f"Q{qn}"
        if slot:
            label += f"({slot})"
        if sub and not is_match_item(item):
            label += f" ({sub})"
        marks = item.get("marks")
        src = item.get("source_type")
        ch = item.get("chapter_or_poem") or item.get("chapter_id")
        space_before = 8.0 if qn != last_q else 4.0
        last_q = qn
        lines.append(
            DocLine(
                f"{label}  [{marks} mark(s)]  source={src}  {ch}",
                "question",
                space_before=space_before,
                space_after=3,
                keep_with_next=True,
                group=f"ak-{label}",
            )
        )
        if is_match_item(item):
            maps = mapping_lines(item)
            if maps:
                for i, row in enumerate(maps):
                    lines.append(
                        DocLine(
                            row,
                            "body",
                            group=f"ak-{label}",
                            space_after=2 if i == len(maps) - 1 else 1,
                        )
                    )
            else:
                lines.append(DocLine(_item_answer_text(item), "body", group=f"ak-{label}", space_after=2))
        else:
            lines.append(DocLine(_item_answer_text(item), "body", group=f"ak-{label}", space_after=2))
        if (item.get("choice_b_answer") or "").strip():
            lines.append(
                DocLine(
                    "OR",
                    "or_line",
                    center=True,
                    space_before=6,
                    space_after=6,
                    keep_with_next=True,
                    group=f"ak-{label}-or",
                )
            )
            lines.append(
                DocLine(
                    item["choice_b_answer"].strip(),
                    "body",
                    group=f"ak-{label}-or",
                )
            )
    return lines


def _txt_meta_block(lines: list[DocLine]) -> list[str]:
    """Render Grade/Subject/Time/Marks as two aligned columns."""
    vals = {ln.text.split(":", 1)[0]: ln.text.split(":", 1)[1].strip() for ln in lines if ":" in ln.text}
    grade = vals.get("Grade", "")
    subject = vals.get("Subject", "")
    time = vals.get("Time", "")
    marks = vals.get("Maximum Marks", "")
    return [
        f"Grade: {grade:<8}  Subject: {subject}",
        f"Time: {time:<9}  Maximum Marks: {marks}",
    ]


def _match_table_txt_lines(table: dict, txt_indent: int = 2) -> list[str]:
    headers = list(table.get("headers") or ["Column A", "Column B"])
    rows = list(table.get("rows") or [])
    col_a = [headers[0]] + [r[0] if r else "" for r in rows]
    col_b = [headers[1] if len(headers) > 1 else "Column B"] + [r[1] if r and len(r) > 1 else "" for r in rows]
    width_a = max((len(str(x)) for x in col_a), default=10)
    width_a = min(max(width_a, 24), 42)
    pad = " " * txt_indent
    out = []
    for a, b in zip(col_a, col_b):
        out.append(f"{pad}{str(a):<{width_a}}  {b}")
    return out


def document_to_txt(doc_lines: list[DocLine]) -> str:
    out: list[str] = []
    i = 0
    while i < len(doc_lines):
        ln = doc_lines[i]
        if ln.group == "header-meta":
            meta = []
            while i < len(doc_lines) and doc_lines[i].group == "header-meta":
                meta.append(doc_lines[i])
                i += 1
            out.append("")
            out.extend(_txt_meta_block(meta))
            out.append("")
            continue
        if ln.style == "section":
            if out and out[-1].strip():
                out.append("")
            out.append(ln.text)
            out.append("")
            i += 1
            continue
        if ln.style == "question":
            if out and out[-1].strip():
                out.append("")
            out.append(ln.text)
            i += 1
            continue
        if ln.table:
            out.extend(_match_table_txt_lines(ln.table, ln.txt_indent or 2))
            i += 1
            continue
        if ln.image_path:
            out.append("  [Picture]")
            i += 1
            continue
        if ln.center:
            if out and out[-1].strip():
                out.append("")
            out.append((ln.text or "").center(60).rstrip())
            if ln.style == "or_line":
                out.append("")
            i += 1
            continue
        if ln.style == "label":
            if out and out[-1].strip():
                out.append("")
            out.append(ln.text)
            i += 1
            continue
        if not ln.text:
            if not out or out[-1] != "":
                out.append("")
            i += 1
            continue
        prefix = " " * ln.txt_indent
        text = prefix + ln.text
        if ln.style == "subq" and out and out[-1].strip().startswith(
            ("      A)", "      B)", "      C)", "      D)", "      (a)", "      (b)")
        ):
            out.append("")
        out.append(text)
        i += 1
    if not out or out[-1] != "":
        out.append("")
    return "\n".join(out)


def render_exam_txt(exam: dict) -> str:
    return document_to_txt(build_exam_document(exam))


def render_answer_key_txt(exam: dict) -> str:
    return document_to_txt(build_answer_key_document(exam))


def _unicode_font_candidates() -> list[Path]:
    return [
        Path(r"C:\Windows\Fonts\Nirmala.ttc"),
        Path(r"C:\Windows\Fonts\Nirmala.ttf"),
        Path(r"C:\Windows\Fonts\nirmala.ttf"),
        Path(r"C:\Windows\Fonts\mangal.ttf"),
        Path(r"C:\Windows\Fonts\arialuni.ttf"),
    ]


def _unicode_fontfile() -> Path | None:
    for p in _unicode_font_candidates():
        if p.exists():
            return p
    return None


def _extract_sfnt_from_ttc(data: bytes, index: int = 0) -> bytes:
    if data[:4] != b"ttcf":
        return data
    num_fonts = struct.unpack_from(">I", data, 8)[0]
    index = max(0, min(index, num_fonts - 1))
    font_offset = struct.unpack_from(">I", data, 12 + 4 * index)[0]
    sfnt_version, num_tables, search_range, entry_selector, range_shift = struct.unpack_from(
        ">4sHHHH", data, font_offset
    )
    records = []
    for i in range(num_tables):
        rec = font_offset + 12 + i * 16
        tag, checksum, offset, length = struct.unpack_from(">4sIII", data, rec)
        records.append((tag, checksum, offset, length))
    header = struct.pack(">4sHHHH", sfnt_version, num_tables, search_range, entry_selector, range_shift)
    out = bytearray(header + b"\x00" * (16 * num_tables))
    for i, (tag, checksum, offset, length) in enumerate(records):
        while len(out) % 4:
            out.append(0)
        new_offset = len(out)
        out.extend(data[offset : offset + length])
        while len(out) % 4:
            out.append(0)
        struct.pack_into(">4sIII", out, 12 + i * 16, tag, checksum, new_offset, length)
    return bytes(out)


def _load_font_buffers() -> tuple[bytes | None, bytes | None]:
    path = _unicode_fontfile()
    if not path:
        return None, None
    data = path.read_bytes()
    regular = _extract_sfnt_from_ttc(data, 0)
    bold = _extract_sfnt_from_ttc(data, 1) if data[:4] == b"ttcf" else regular
    return regular, bold


class _FontSet:
    def __init__(self):
        import fitz

        reg_buf, bold_buf = _load_font_buffers()
        self.regular = fitz.Font(fontbuffer=reg_buf) if reg_buf else fitz.Font("helv")
        self.bold = fitz.Font(fontbuffer=bold_buf) if bold_buf else fitz.Font("helv")
        if self.bold.name == self.regular.name and reg_buf and bold_buf and bold_buf == reg_buf:
            # last-resort: still usable; caller may double-stroke
            pass

    def font(self, bold: bool):
        return self.bold if bold else self.regular

    def width(self, text: str, size: float, bold: bool) -> float:
        try:
            return self.font(bold).text_length(text, fontsize=size)
        except Exception:
            return len(text) * size * 0.5


def _wrap_text(text: str, fonts: _FontSet, size: float, bold: bool, max_width: float) -> list[str]:
    if not text:
        return [""]
    words = text.split(" ")
    rows: list[str] = []
    cur = ""
    for w in words:
        trial = w if not cur else f"{cur} {w}"
        if fonts.width(trial, size, bold) <= max_width or not cur:
            cur = trial
        else:
            rows.append(cur)
            cur = w
    rows.append(cur)
    return rows


def _line_metrics(style: str) -> tuple[float, bool, float]:
    spec = STYLES.get(style, STYLES["body"])
    size = float(spec["size"])
    bold = bool(spec["bold"])
    leading = size * float(spec["leading"])
    return size, bold, leading


@dataclass
class _Frag:
    text: str
    style: str
    indent: float
    center: bool
    space_before: float
    space_after: float
    group: str
    size: float
    bold: bool
    leading: float
    keep_with_next: bool = False
    image_path: str = ""
    image_height: float = 0.0
    table: dict | None = None
    table_height: float = 0.0

    @property
    def height(self) -> float:
        if self.table:
            return self.space_before + self.table_height + self.space_after
        if self.image_path:
            return self.space_before + self.image_height + self.space_after
        return self.space_before + self.leading + self.space_after


def _match_table_col_widths(table: dict, usable: float) -> list[float]:
    weights = list(table.get("col_weights") or [0.62, 0.38])
    if len(weights) < 2:
        weights = [0.62, 0.38]
    total = sum(weights[:2]) or 1.0
    return [usable * (weights[0] / total), usable * (weights[1] / total)]


def _match_table_row_lines(table: dict, fonts: _FontSet, usable: float, size: float) -> list[list[list[str]]]:
    widths = _match_table_col_widths(table, usable)
    pad = 6.0
    headers = list(table.get("headers") or ["Column A", "Column B"])
    body = list(table.get("rows") or [])
    all_rows = [headers[:2]] + [list(r[:2]) + [""] * max(0, 2 - len(r)) for r in body]
    wrapped_rows = []
    for row in all_rows:
        wrapped_rows.append(
            [
                _wrap_text(str(row[0] if len(row) > 0 else ""), fonts, size, False, max(40.0, widths[0] - 2 * pad)),
                _wrap_text(str(row[1] if len(row) > 1 else ""), fonts, size, False, max(40.0, widths[1] - 2 * pad)),
            ]
        )
    return wrapped_rows


def _match_table_height(table: dict, fonts: _FontSet, usable: float, size: float, leading: float) -> float:
    pad = 6.0
    rows = _match_table_row_lines(table, fonts, usable, size)
    return sum(max(len(cells[0]), len(cells[1]), 1) * leading + 2 * pad for cells in rows)


def _flatten_frags(doc_lines: list[DocLine], fonts: _FontSet, max_width: float) -> list[_Frag]:
    frags: list[_Frag] = []
    for ln in doc_lines:
        size, bold, leading = _line_metrics(ln.style)
        if ln.table:
            usable = max(120.0, max_width - ln.indent)
            table_h = _match_table_height(ln.table, fonts, usable, size, leading)
            frags.append(
                _Frag(
                    text="",
                    style=ln.style,
                    indent=ln.indent,
                    center=False,
                    space_before=ln.space_before,
                    space_after=ln.space_after,
                    group=ln.group,
                    size=size,
                    bold=bold,
                    leading=leading,
                    keep_with_next=ln.keep_with_next,
                    table=ln.table,
                    table_height=table_h,
                )
            )
            continue
        if ln.image_path:
            frags.append(
                _Frag(
                    text="",
                    style=ln.style,
                    indent=ln.indent,
                    center=ln.center,
                    space_before=ln.space_before,
                    space_after=ln.space_after,
                    group=ln.group,
                    size=size,
                    bold=bold,
                    leading=leading,
                    keep_with_next=ln.keep_with_next,
                    image_path=ln.image_path,
                    image_height=ln.image_height or 168.0,
                )
            )
            continue
        usable = max(80.0, max_width - ln.indent)
        wrapped = _wrap_text(ln.text, fonts, size, bold, usable) if ln.text else [""]
        for i, row in enumerate(wrapped):
            frags.append(
                _Frag(
                    text=row,
                    style=ln.style,
                    indent=ln.indent,
                    center=ln.center,
                    space_before=ln.space_before if i == 0 else 0.0,
                    space_after=ln.space_after if i == len(wrapped) - 1 else 0.0,
                    group=ln.group,
                    size=size,
                    bold=bold,
                    leading=leading,
                    keep_with_next=ln.keep_with_next if i == len(wrapped) - 1 else True,
                )
            )
    return frags


def _group_height(frags: list[_Frag], start: int) -> tuple[float, int]:
    if start >= len(frags):
        return 0.0, start
    gid = frags[start].group
    total = 0.0
    i = start
    if gid:
        while i < len(frags) and frags[i].group == gid:
            total += frags[i].height
            i += 1
    else:
        total = frags[start].height
        i = start + 1
    # Also keep a heading/OR with the next content block when flagged.
    j = start
    keep_total = 0.0
    while j < len(frags):
        keep_total += frags[j].height
        if not frags[j].keep_with_next or j + 1 >= len(frags):
            j += 1
            break
        nxt = frags[j + 1].height
        if keep_total + nxt > KEEP_MAX_PT:
            j += 1
            break
        j += 1
    if keep_total > total:
        return keep_total, max(i, j)
    return total, i


def _draw_match_table(page, fonts: _FontSet, x0: float, y0: float, usable: float, table: dict, size: float, leading: float) -> float:
    import fitz

    pad = 6.0
    widths = _match_table_col_widths(table, usable)
    wrapped_rows = _match_table_row_lines(table, fonts, usable, size)
    writer = fitz.TextWriter(page.rect)
    y = y0
    for r_i, cells in enumerate(wrapped_rows):
        row_h = max(len(cells[0]), len(cells[1]), 1) * leading + 2 * pad
        x = x0
        for c_i, lines in enumerate(cells):
            w = widths[c_i]
            rect = fitz.Rect(x, y, x + w, y + row_h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.7)
            ty = y + pad + size * 0.85
            font = fonts.font(r_i == 0)
            for line in lines:
                writer.append((x + pad, ty), line, font=font, fontsize=size)
                ty += leading
            x += w
        y += row_h
    writer.write_text(page)
    return y


def render_styled_pdf(doc_lines: list[DocLine], dest: Path) -> None:
    import fitz

    fonts = _FontSet()
    doc = fitz.open()
    doc.new_page()
    width, height = doc[0].rect.width, doc[0].rect.height
    max_width = width - 2 * MARGIN_PT
    footer_y = height - 28
    content_bottom = height - MARGIN_PT
    y = MARGIN_PT + 12
    writer = fitz.TextWriter(doc[0].rect)

    def new_page():
        nonlocal y, writer
        writer.write_text(doc[-1])
        doc.new_page()
        writer = fitz.TextWriter(doc[-1].rect)
        y = MARGIN_PT + 12

    def draw_meta_grid(start: int) -> int:
        nonlocal y
        block = []
        j = start
        while j < len(frags) and frags[j].group == "header-meta":
            if frags[j].text:
                block.append(frags[j])
            j += 1
        if y + 36 > content_bottom and y > MARGIN_PT + 40:
            new_page()
        col2 = width / 2
        size, bold, leading = _line_metrics("meta")
        pairs = [("Grade", "Subject"), ("Time", "Maximum Marks")]
        by_key = {}
        for frag in block:
            if ":" in frag.text:
                k, v = frag.text.split(":", 1)
                by_key[k.strip()] = f"{k.strip()}: {v.strip()}"
        y += 4
        for left_k, right_k in pairs:
            left = by_key.get(left_k, "")
            right = by_key.get(right_k, "")
            if left:
                writer.append((MARGIN_PT, y + size * 0.85), left, font=fonts.font(False), fontsize=size)
            if right:
                writer.append((col2, y + size * 0.85), right, font=fonts.font(False), fontsize=size)
            y += leading + 3
        y += 8
        return j

    frags = _flatten_frags(doc_lines, fonts, max_width)
    i = 0
    while i < len(frags):
        if frags[i].group == "header-meta":
            i = draw_meta_grid(i)
            continue
        g_h, g_end = _group_height(frags, i)
        remaining = content_bottom - y
        if (
            frags[i].group
            and g_h <= KEEP_MAX_PT
            and g_h > remaining
            and y > MARGIN_PT + 40
        ):
            new_page()
        frag = frags[i]
        needed = frag.height
        if y + needed > content_bottom and y > MARGIN_PT + 40:
            new_page()
        if frag.table:
            y += frag.space_before
            writer.write_text(doc[-1])
            y = _draw_match_table(
                doc[-1],
                fonts,
                MARGIN_PT + frag.indent,
                y,
                max(120.0, max_width - frag.indent),
                frag.table,
                frag.size,
                frag.leading,
            )
            writer = fitz.TextWriter(doc[-1].rect)
            y += frag.space_after
            i += 1
            continue
        if frag.image_path:
            y += frag.space_before
            img_h = frag.image_height or 168.0
            img_w = min(max_width - frag.indent, img_h * 1.7)
            writer.write_text(doc[-1])
            rect = fitz.Rect(
                MARGIN_PT + frag.indent,
                y,
                MARGIN_PT + frag.indent + img_w,
                y + img_h,
            )
            try:
                doc[-1].insert_image(rect, filename=frag.image_path, keep_proportion=True)
            except Exception as e:
                log.error("Failed to embed picture image: %s", e)
            writer = fitz.TextWriter(doc[-1].rect)
            y += img_h + frag.space_after
            i += 1
            continue
        y += frag.space_before
        text = frag.text
        if text:
            tw = fonts.width(text, frag.size, frag.bold)
            if frag.center:
                x = max(MARGIN_PT, (width - tw) / 2)
            else:
                x = MARGIN_PT + frag.indent
            writer.append(
                (x, y + frag.size * 0.85),
                text,
                font=fonts.font(frag.bold),
                fontsize=frag.size,
            )
            if frag.bold and fonts.bold.name == fonts.regular.name:
                writer.append(
                    (x + 0.35, y + frag.size * 0.85),
                    text,
                    font=fonts.font(False),
                    fontsize=frag.size,
                )
        y += frag.leading + frag.space_after
        i += 1

    writer.write_text(doc[-1])
    total = doc.page_count
    footer_size = STYLES["footer"]["size"]
    for idx in range(total):
        pg = doc[idx]
        label = f"Page {idx + 1} of {total}"
        tw = fonts.width(label, footer_size, False)
        fw = fitz.TextWriter(pg.rect)
        fw.append(((width - tw) / 2, footer_y), label, font=fonts.regular, fontsize=footer_size)
        fw.write_text(pg)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)
    doc.close()
    log.info("Wrote exam PDF %s", dest)




def _docx_set_keep_with_next(paragraph, value: bool) -> None:
    paragraph.paragraph_format.keep_with_next = bool(value)


def _docx_set_cell_shading(cell, fill: str = "E7E6E6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _docx_set_run_font(run, font_name: str = "Noto Sans") -> None:
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), font_name)


def _docx_add_text(paragraph, text: str, *, size: float, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    _docx_set_run_font(run)
    run.font.size = Pt(size)
    run.bold = bool(bold)


def render_styled_docx(doc_lines: list[DocLine], dest: Path) -> None:
    """Write a Word document from the same canonical DocLine stream used by PDF/TXT."""
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Noto Sans"
    normal.font.size = Pt(STYLES["body"]["size"])

    for ln in doc_lines:
        if ln.table:
            rows = list(ln.table.get("rows") or [])
            headers = list(ln.table.get("headers") or ["Column A", "Column B"])
            table = doc.add_table(rows=1, cols=max(2, len(headers)))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            hdr = table.rows[0].cells
            for idx, text in enumerate(headers[: len(hdr)]):
                hdr[idx].text = ""
                hdr[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = hdr[idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _docx_add_text(p, str(text), size=11.0, bold=True)
                _docx_set_cell_shading(hdr[idx])
            for row in rows:
                cells = table.add_row().cells
                values = list(row) if isinstance(row, (list, tuple)) else [str(row)]
                for idx, cell in enumerate(cells):
                    cell.text = ""
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    value = values[idx] if idx < len(values) else ""
                    _docx_add_text(p, str(value), size=10.5, bold=False)
            tail = doc.add_paragraph()
            tail.paragraph_format.space_after = Pt(max(ln.space_after, 2.0))
            continue

        p = doc.add_paragraph()
        fmt = p.paragraph_format
        fmt.left_indent = Pt(max(ln.indent, 0.0)) if ln.indent else None
        fmt.space_before = Pt(max(ln.space_before, 0.0))
        fmt.space_after = Pt(max(ln.space_after, 0.0))
        fmt.line_spacing = STYLES.get(ln.style, STYLES["body"]).get("leading", 1.15)
        _docx_set_keep_with_next(p, ln.keep_with_next)

        if ln.center or ln.style in {"school", "title"}:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if ln.image_path:
            image = Path(ln.image_path)
            if image.exists():
                run = p.add_run()
                height = Inches((ln.image_height or 168.0) / 72.0)
                run.add_picture(str(image), height=height)
            else:
                _docx_add_text(p, "[Picture]", size=STYLES["body"]["size"], bold=False)
            continue

        style = STYLES.get(ln.style, STYLES["body"])
        _docx_add_text(
            p,
            ln.text or "",
            size=float(style.get("size", STYLES["body"]["size"])),
            bold=bool(style.get("bold", False)),
        )

    doc.save(dest)
    log.info("Wrote exam DOCX %s", dest)


def render_exam_docx(exam: dict, dest: Path) -> None:
    """Write the student exam Word document using the same content/order as PDF/TXT."""
    render_styled_docx(build_exam_document(exam), dest)

def _plain_text_document(text: str) -> list[DocLine]:
    lines: list[DocLine] = []
    for raw in text.splitlines():
        s = raw.rstrip()
        if not s.strip():
            lines.append(DocLine("", space_after=4))
            continue
        stripped = s.strip()
        if stripped.upper() in {"ANSWER KEY", "SYLLABUS", "INSTRUCTIONS"}:
            style = "school" if stripped.upper() == "ANSWER KEY" else "label"
            lines.append(DocLine(stripped, style, space_before=6, space_after=4))
        elif stripped.startswith("SECTION "):
            lines.append(DocLine(stripped, "section", space_before=8, space_after=6))
        elif re.match(r"^Q\d+", stripped):
            lines.append(DocLine(stripped, "question", space_before=8, space_after=3, keep_with_next=True))
        elif stripped == "OR":
            lines.append(DocLine("OR", "or_line", center=True, space_before=8, space_after=8, keep_with_next=True))
        elif re.match(r"^section\s+[a-z]\b", stripped, re.I):
            lines.append(DocLine(stripped, "section", space_before=8, space_after=6))
        else:
            indent = 0.0
            txt_indent = len(s) - len(s.lstrip(" "))
            if txt_indent >= 6:
                indent = INDENT_OPTION
            elif txt_indent >= 2:
                indent = INDENT_SUBQ
            lines.append(DocLine(stripped, "body", indent=indent, txt_indent=txt_indent))
    return lines


def render_exam_pdf(exam: dict, dest: Path) -> None:
    """Write the student exam PDF using the same content/order as TXT."""
    render_styled_pdf(build_exam_document(exam), dest)


def render_answer_key_pdf(exam: dict, dest: Path) -> None:
    render_styled_pdf(build_answer_key_document(exam), dest)


def render_text_pdf(text: str, dest: Path) -> None:
    render_styled_pdf(_plain_text_document(text), dest)
